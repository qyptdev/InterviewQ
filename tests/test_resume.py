"""Tests for resume upload and question generation."""

import pytest
from unittest.mock import patch
from fastapi.testclient import TestClient
from app.main import app
from app.database import init_db


@pytest.fixture
def client():
    """Create test client."""
    init_db()
    with TestClient(app) as c:
        yield c


class TestDocumentParser:
    """Tests for document parsing functions."""

    def test_parse_text(self):
        """Test plain text parsing."""
        from app.services.document_parser import parse_text

        result = parse_text("  Hello World  ")
        assert result == "Hello World"

    def test_parse_markdown(self):
        """Test markdown parsing strips formatting."""
        from app.services.document_parser import parse_markdown

        md = "# Header\n**Bold** text\n*Italic* word"
        result = parse_markdown(md)
        assert "Header" in result
        assert "Bold" in result
        assert "**" not in result

    def test_parse_pdf_bytes(self):
        """Test PDF parsing with valid PDF bytes."""
        import fitz

        # Create a minimal valid PDF in memory
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Test Resume Content")
        pdf_bytes = doc.tobytes()
        doc.close()

        from app.services.document_parser import parse_pdf

        result = parse_pdf(pdf_bytes)
        assert "Test Resume Content" in result

    def test_parse_docx_bytes(self):
        """Test DOCX parsing with valid DOCX bytes."""
        from docx import Document
        import io

        doc = Document()
        doc.add_paragraph("Resume Summary")
        doc.add_paragraph("Work Experience at Company A")
        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()

        from app.services.document_parser import parse_docx

        result = parse_docx(docx_bytes)
        assert "Resume Summary" in result
        assert "Work Experience at Company A" in result

    def test_parse_document_binary_unsupported(self):
        """Test binary parser rejects unsupported types."""
        from app.services.document_parser import parse_document_binary

        with pytest.raises(ValueError, match="不支持的二进制文件格式"):
            parse_document_binary(b"content", "xyz")

    def test_extract_resume_sections(self):
        """Test resume section extraction with expanded fields."""
        from app.services.document_parser import extract_resume_sections

        text = """专业技能
Python, FastAPI, SQLite

工作经验
3年Web后端开发

教育经历
计算机科学学士

项目经历
AI面试系统开发"""

        sections = extract_resume_sections(text)
        assert "Python" in sections["skills"]
        assert "3年" in sections["experience"]
        assert "计算机" in sections["education"]
        assert "AI面试" in sections["projects"]
        # Check new fields exist
        assert "basic_info" in sections
        assert "summary" in sections

    def test_extract_resume_sections_all_keys_present(self):
        """Test that all 7 section keys are always present."""
        from app.services.document_parser import extract_resume_sections

        sections = extract_resume_sections("some random text without headers")
        expected_keys = {"basic_info", "summary", "education", "skills", "experience", "projects", "other"}
        assert set(sections.keys()) == expected_keys


class TestNodesValidate:
    """Tests for validation node type safety."""

    def test_validate_list_title(self):
        """Test validation handles list type titles from LLM."""
        from app.core.nodes_validate import validate_question_batch
        from app.core.state_models import QuestionBatch

        batch = QuestionBatch(
            batch_index=0,
            questions=[
                {"title": ["What", "is", "Python?"], "expected_answer": "A language"},
                {"title": "Explain REST", "expected_answer": "An architecture"},
            ],
        )
        result = validate_question_batch(batch)
        # Should not crash, title should be joined into string
        assert isinstance(result, object)
        # The list title should have been corrected to a string
        assert batch.questions[0]["title"] == "What is Python?"

    def test_validate_list_expected_answer(self):
        """Test validation handles list type expected_answer from LLM."""
        from app.core.nodes_validate import validate_question_batch
        from app.core.state_models import QuestionBatch

        batch = QuestionBatch(
            batch_index=0,
            questions=[
                {
                    "title": "What is Python programming?",
                    "expected_answer": ["Python", "is", "a", "language"],
                },
            ],
        )
        result = validate_question_batch(batch)
        assert batch.questions[0]["expected_answer"] == "Python is a language"

    def test_deduplicate_list_titles(self):
        """Test deduplication handles list type titles."""
        from app.core.nodes_validate import deduplicate_questions

        questions = [
            {"title": ["What", "is", "Python?"]},
            {"title": "What is Python?"},
            {"title": "Explain Docker"},
        ]
        result = deduplicate_questions(questions)
        # First two are duplicates after normalization
        assert len(result) == 2

    def test_validate_empty_title_from_list(self):
        """Test validation catches empty title from empty list."""
        from app.core.nodes_validate import validate_question_batch
        from app.core.state_models import QuestionBatch

        batch = QuestionBatch(
            batch_index=0,
            questions=[
                {"title": [], "expected_answer": ""},
            ],
        )
        result = validate_question_batch(batch)
        assert not result.is_valid
        assert any("标题为空" in issue for issue in result.issues)


class TestAIExtractResumeSections:
    """Tests for AI resume section extraction."""

    @patch("app.llm.client.get_llm_router")
    def test_ai_extract_success(self, mock_get_router):
        """Test AI extraction returns parsed sections with 7 fields on success."""
        import asyncio
        import json
        from unittest.mock import AsyncMock
        from app.services.document_parser import ai_extract_resume_sections

        mock_router = AsyncMock()
        mock_router.generate_with_fallback = AsyncMock(
            return_value=json.dumps({
                "basic_info": "John Doe, 123-456-7890",
                "summary": "Experienced developer",
                "experience": "5 years at Google",
                "education": "MIT CS",
                "skills": "Python, Go",
                "projects": "Built search engine",
                "other": "AWS certified",
            })
        )
        mock_get_router.return_value = mock_router

        result = asyncio.get_event_loop().run_until_complete(
            ai_extract_resume_sections("some resume text")
        )
        assert isinstance(result, dict)
        assert result["summary"] == "Experienced developer"
        assert result["skills"] == "Python, Go"
        assert result["basic_info"] == "John Doe, 123-456-7890"
        assert result["projects"] == "Built search engine"

    @patch("app.llm.client.get_llm_router")
    def test_ai_extract_fallback_on_failure(self, mock_get_router):
        """Test AI extraction falls back to fixed parsing on LLM failure."""
        import asyncio
        from unittest.mock import AsyncMock
        from app.services.document_parser import ai_extract_resume_sections

        mock_router = AsyncMock()
        mock_router.generate_with_fallback = AsyncMock(side_effect=Exception("LLM down"))
        mock_get_router.return_value = mock_router

        text = "专业技能\nPython, FastAPI\n\n工作经验\n3年开发"
        result = asyncio.get_event_loop().run_until_complete(
            ai_extract_resume_sections(text)
        )
        # Should return dict even on failure (fallback to fixed parsing)
        assert isinstance(result, dict)
        assert "summary" in result
        assert "skills" in result
        assert "experience" in result
        assert "basic_info" in result
        assert "projects" in result


class TestQuestionGenerator:
    """Tests for question generator service."""

    def test_parse_uploaded_file_txt(self):
        """Test parsing uploaded text file."""
        import asyncio
        from app.services.question_generator import parse_uploaded_file

        content = "Hello, this is a test resume.".encode("utf-8")
        result = asyncio.get_event_loop().run_until_complete(
            parse_uploaded_file("resume.txt", content)
        )
        assert "Hello" in result

    def test_parse_uploaded_file_empty_name(self):
        """Test parsing rejects empty filename."""
        import asyncio
        from app.services.question_generator import parse_uploaded_file

        with pytest.raises(ValueError, match="文件名不能为空"):
            asyncio.get_event_loop().run_until_complete(
                parse_uploaded_file("", b"content")
            )

    def test_parse_uploaded_file_unsupported_ext(self):
        """Test parsing rejects unsupported extension."""
        import asyncio
        from app.services.question_generator import parse_uploaded_file

        with pytest.raises(ValueError, match="不支持的文件格式"):
            asyncio.get_event_loop().run_until_complete(
                parse_uploaded_file("resume.exe", b"content")
            )

    def test_save_questions_to_bank(self):
        """Test saving questions creates bank and links questions."""
        from app.services.question_generator import save_questions_to_bank

        questions = [
            {
                "title": "What is Python?",
                "category": "技术",
                "difficulty": "easy",
                "expected_answer": "A programming language.",
            },
            {
                "title": "Explain REST API",
                "category": "技术",
                "difficulty": "medium",
                "expected_answer": "REST is an architectural style.",
            },
        ]

        result = save_questions_to_bank(questions, bank_name="Test Bank")
        assert result["bank"]["name"] == "Test Bank"
        assert result["questions_saved"] == 2
        assert len(result["question_ids"]) == 2

    def test_save_questions_auto_bank_name(self):
        """Test auto-generated bank name when not provided."""
        from app.services.question_generator import save_questions_to_bank

        questions = [{"title": "Q1", "category": "技术", "difficulty": "easy", "expected_answer": ""}]
        result = save_questions_to_bank(questions)
        assert "简历面试题-" in result["bank"]["name"]

    def test_save_questions_empty_list(self):
        """Test saving empty questions list raises error."""
        from app.services.question_generator import save_questions_to_bank

        with pytest.raises(ValueError, match="没有可保存的题目"):
            save_questions_to_bank([])

    @patch("app.services.question_generator.plan_and_execute")
    def test_generate_questions_from_resume(self, mock_plan):
        """Test question generation calls plan_and_execute with job_title."""
        import asyncio

        mock_plan.return_value = [
            {"title": "Q1", "expected_answer": "A1"},
            {"title": "Q2", "expected_answer": "A2"},
        ]

        from app.services.question_generator import generate_questions_from_resume

        result = asyncio.get_event_loop().run_until_complete(
            generate_questions_from_resume("resume text", "jd text", "产品经理", 5)
        )
        assert len(result) == 2
        mock_plan.assert_called_once_with(
            resume_text="resume text",
            jd_text="jd text",
            job_title="产品经理",
            question_count=5,
            generation_mode="standard",
        )

    def test_generate_questions_empty_resume(self):
        """Test generation rejects empty resume."""
        import asyncio
        from app.services.question_generator import generate_questions_from_resume

        with pytest.raises(ValueError, match="简历内容不能为空"):
            asyncio.get_event_loop().run_until_complete(
                generate_questions_from_resume("", "", "", 5)
            )


class TestResumeAPI:
    """Tests for resume API endpoints."""

    def test_upload_resume_txt(self, client):
        """Test uploading a text resume."""
        files = {"file": ("resume.txt", b"My resume content", "text/plain")}
        response = client.post("/api/resume/upload", files=files)
        assert response.status_code == 200
        data = response.json()
        assert data["filename"] == "resume.txt"
        assert "sections" in data
        assert data["parse_method"] == "fixed"
        # Check expanded sections
        assert "basic_info" in data["sections"]
        assert "projects" in data["sections"]

    def test_upload_resume_unsupported(self, client):
        """Test uploading unsupported file type."""
        files = {"file": ("resume.exe", b"content", "application/octet-stream")}
        response = client.post("/api/resume/upload", files=files)
        assert response.status_code == 400

    def test_upload_resume_empty_filename(self, client):
        """Test uploading file with empty filename."""
        files = {"file": ("", b"content", "text/plain")}
        response = client.post("/api/resume/upload", files=files)
        # FastAPI returns 422 for validation errors on empty filename
        assert response.status_code in (400, 422)

    def test_ai_parse_resume_txt(self, client):
        """Test AI parse endpoint with text file (falls back to fixed parsing)."""
        files = {"file": ("resume.txt", b"Skills\nPython\nExperience\n3 years", "text/plain")}
        response = client.post("/api/resume/ai-parse", files=files)
        # Should succeed (AI parsing may fail and fallback to fixed)
        assert response.status_code == 200
        data = response.json()
        assert data["filename"] == "resume.txt"
        assert "sections" in data
        assert data["parse_method"] == "ai"
        # Check expanded sections
        assert "basic_info" in data["sections"]
        assert "projects" in data["sections"]

    def test_ai_parse_resume_unsupported(self, client):
        """Test AI parse rejects unsupported file type."""
        files = {"file": ("resume.exe", b"content", "application/octet-stream")}
        response = client.post("/api/resume/ai-parse", files=files)
        assert response.status_code == 400

    @patch("app.services.question_generator.plan_and_execute")
    def test_generate_from_resume(self, mock_plan, client):
        """Test full generate pipeline via API with job_title."""
        mock_plan.return_value = [
            {
                "title": "Explain your Python experience",
                "category": "技术",
                "difficulty": "medium",
                "expected_answer": "Should discuss projects and libraries used.",
            },
            {
                "title": "What is a REST API?",
                "category": "技术",
                "difficulty": "easy",
                "expected_answer": "An architectural style for web services.",
            },
        ]

        files = {"file": ("resume.txt", b"5 years Python developer", "text/plain")}
        data = {"question_count": 2, "jd_text": "Python developer position", "job_title": "Python工程师"}
        response = client.post("/api/resume/generate", files=files, data=data)

        assert response.status_code == 200
        result = response.json()
        assert result["questions_saved"] == 2
        assert result["bank_id"] > 0
        assert "resume_sections" in result
        # Verify job_title was passed through
        call_kwargs = mock_plan.call_args[1]
        assert call_kwargs["job_title"] == "Python工程师"

    @patch("app.services.question_generator.plan_and_execute")
    def test_generate_with_edited_sections(self, mock_plan, client):
        """Test generate pipeline with user-edited sections (7 fields)."""
        import json

        mock_plan.return_value = [
            {
                "title": "Tell me about your Python projects",
                "category": "技术",
                "difficulty": "medium",
                "expected_answer": "Discuss past projects.",
            },
        ]

        edited = json.dumps({
            "basic_info": "John Doe",
            "summary": "Custom summary",
            "experience": "Custom experience",
            "education": "Custom education",
            "skills": "Python, FastAPI",
            "projects": "AI Chatbot",
            "other": "",
        })

        files = {"file": ("resume.txt", b"5 years Python developer", "text/plain")}
        data = {"question_count": 1, "edited_sections": edited, "job_title": "后端工程师"}
        response = client.post("/api/resume/generate", files=files, data=data)

        assert response.status_code == 200
        result = response.json()
        assert result["questions_saved"] == 1
        # The sections should be the user-edited ones
        assert result["resume_sections"]["summary"] == "Custom summary"
        assert result["resume_sections"]["skills"] == "Python, FastAPI"
        assert result["resume_sections"]["basic_info"] == "John Doe"
        assert result["resume_sections"]["projects"] == "AI Chatbot"
